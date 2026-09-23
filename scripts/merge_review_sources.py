"""Append two review sources without altering values or dropping duplicates."""

import argparse
import os
from pathlib import Path
import sys
import tempfile
from zipfile import BadZipFile

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils.exceptions import InvalidFileException

from inspect_reviews import COLUMNS, PROJECT_ROOT, is_blank, markdown_value


PRODUCT_KEY = ["platform", "product_id"]
REVIEW_KEY = [*PRODUCT_KEY, "review_id"]


def read_source(path, sheet_name):
    if not path.is_file():
        raise ValueError(f"입력 파일이 없습니다: {path}")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"입력 파일에 '{sheet_name}' 시트가 없습니다: {path}")
        sheet = workbook[sheet_name]
        header_row = None
        positions = None
        header_count = 0
        # First pass finds the last full, unambiguous header, never a fixed row.
        for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            labels = [value.strip() if isinstance(value, str) else value for value in row]
            if all(labels.count(column) == 1 for column in COLUMNS):
                header_row = number
                positions = [labels.index(column) for column in COLUMNS]
                header_count += 1
        if header_row is None:
            raise ValueError(f"'{sheet_name}' 시트에서 공통 컬럼 11개가 각각 한 번 있는 header를 찾지 못했습니다.")
        records = []
        skipped = 0
        # Examples and template instructions above the last header are excluded.
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            record = tuple(row[index] if index < len(row) else None for index in positions)
            if all(is_blank(value) for value in record):
                skipped += 1
                continue
            records.append(record)
        return records, {"path": path, "sheet": sheet_name, "header": header_row,
                         "headers": header_count, "skipped": skipped}
    finally:
        workbook.close()


def diagnostics(base_records, new_records):
    # Object dtype is for statistics only. Export uses original record tuples.
    frame = pd.DataFrame([*base_records, *new_records], columns=COLUMNS, dtype=object)
    blanks = frame.apply(lambda column: column.map(is_blank)).astype(bool)
    valid_products = frame.loc[~blanks[PRODUCT_KEY].any(axis=1)]
    valid_keys = frame.loc[~blanks[REVIEW_KEY].any(axis=1), REVIEW_KEY]
    sizes = valid_keys.groupby(REVIEW_KEY, sort=False).size()
    base_keys = valid_keys.loc[valid_keys.index < len(base_records)].drop_duplicates()
    new_keys = valid_keys.loc[valid_keys.index >= len(base_records)].drop_duplicates()
    # Intersect only complete composite keys; repeated rows still remain in output.
    combined_keys = pd.concat([base_keys, new_keys], ignore_index=True)
    overlap = int(combined_keys.duplicated(subset=REVIEW_KEY).sum())
    stats = [
        ("기존 데이터 리뷰 수", len(base_records)),
        ("신규 데이터 리뷰 수", len(new_records)),
        ("병합 전체 리뷰 수", len(frame)),
        ("고유 상품 수 ((platform, product_id), 빈값 제외)", len(valid_products.drop_duplicates(subset=PRODUCT_KEY))),
        ("content 빈값 수", int(blanks["content"].sum())),
        ("복합 리뷰키 중복 그룹 수", int((sizes > 1).sum())),
        ("복합 리뷰키 중복 참여 행 수 (최초 행 포함)", int(sizes[sizes > 1].sum())),
        ("기존↔신규 사이에 겹치는 복합 리뷰키 수", overlap),
        ("복합 리뷰키 구성값 누락으로 중복 검사에서 제외된 행 수", int(blanks[REVIEW_KEY].any(axis=1).sum())),
    ]
    distributions = {}
    for column in ("platform", "rating"):
        entries = list(frame.loc[~blanks[column], column].value_counts(sort=False).items())
        entries.append(("[빈값 합계]", int(blanks[column].sum())))
        distributions[column] = entries
    return stats, distributions


def display_path(path):
    try:
        return path.relative_to(PROJECT_ROOT).as_posix()
    except ValueError:
        return str(path)


def report_table(title, entries):
    return [f"## {title}", "", "| 항목 | 개수 |", "| --- | ---: |",
            *(f"| {markdown_value(name)} | {int(count)} |" for name, count in entries), ""]


def build_report(base_info, new_info, stats, distributions):
    lines = ["# 리뷰 source 병합 요약", ""]
    for name, info in (("기존", base_info), ("신규", new_info)):
        lines += [
            f"- {name} 입력: {markdown_value(display_path(info['path']))}",
            f"- {name} 시트: {markdown_value(info['sheet'])}",
            f"- {name} 실제 header 위치: Excel {info['header']}행 (유효 header {info['headers']}개 발견)",
            f"- {name} 공통 컬럼 전체 빈 행 제외: {info['skipped']}행",
        ]
    lines += ["", "## 처리 원칙", "",
              "- 각 시트에서 공통 컬럼 11개를 각각 한 번 포함한 마지막 header 이후만 읽습니다. 그 이전의 예시/안내 영역은 포함하지 않습니다.",
              "- 기존 데이터 다음에 신규 데이터를 이어 붙이며 각 source의 행 순서를 유지합니다.",
              "- 공통 컬럼이 모두 빈 행만 제외합니다. None/NaN, 빈 문자열 및 공백만 있는 문자열을 빈값으로 검사하며 원본 값을 대체하지 않습니다.",
              "- 출력 Reviews 시트에는 공통 Review 컬럼 11개만 기록합니다. content, ID, rating 등 원본 셀 값을 정규화하거나 보완하지 않습니다.",
              "- 고유 상품은 (platform, product_id), 리뷰 중복은 (platform, product_id, review_id) 기준입니다. 각 키의 모든 값이 존재할 때만 계산합니다.",
              "- 기존↔신규 중복 수는 두 source에 모두 존재하는 서로 다른 복합 리뷰키 수이며 pair 수가 아닙니다.",
              "- 중복 행도 모두 유지합니다. 통계는 진단용이며 자동 삭제, dedup, 라벨링을 하지 않습니다.",
              "- ID 타입을 강제 통일하지 않습니다. 통계 비교에서 숫자 int/float는 같게 취급될 수 있지만 숫자와 문자열은 구분합니다.",
              "- 수식 셀은 계산하지 않고 원본 수식 문자열을 보존합니다. Excel 출력에서도 문자열로 기록합니다.",
              "- Excel 저장 형식상 int/float의 세부 구분 등 Python 타입의 완전한 왕복 보존은 보장되지 않습니다.",
              "- 보고서에는 리뷰 본문을 싣지 않습니다.", ""]
    lines += report_table("병합 및 중복 통계", stats)
    lines += report_table("플랫폼별 리뷰 수", distributions["platform"])
    lines += report_table("rating 분포", distributions["rating"])
    return "\n".join(lines)


def write_workbook(stream, base_records, new_records):
    if len(base_records) + len(new_records) > 1_048_575:
        raise ValueError("병합 리뷰 수가 Excel 단일 시트 행 한도를 초과했습니다.")
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Reviews")
    sheet.append(list(COLUMNS))
    try:
        for records in (base_records, new_records):
            for record in records:
                cells = []
                for value in record:
                    cell = WriteOnlyCell(sheet, value=value)
                    if isinstance(value, str):
                        cell.data_type = "s"
                    cells.append(cell)
                sheet.append(cells)
        workbook.save(stream)
    finally:
        workbook.close()


def publish_no_replace(temporary_path, destination):
    if os.name == "nt":
        # Windows rename is atomic and refuses an existing destination.
        temporary_path.rename(destination)
    else:
        # POSIX rename overwrites; an exclusive hard link prevents that race.
        os.link(temporary_path, destination)


def main():
    parser = argparse.ArgumentParser(description="두 리뷰 source를 원본 변경 및 중복 삭제 없이 병합합니다.")
    parser.add_argument("--base", type=Path, default=PROJECT_ROOT / "raw" / "폼_Review_data.xlsx")
    parser.add_argument("--base-sheet", default="Reviews")
    parser.add_argument("--new", type=Path, default=PROJECT_ROOT / "raw" / "폼_Review_data_0919.xlsx")
    parser.add_argument("--new-sheet", default="Review0919")
    args = parser.parse_args()
    base_path = args.base.expanduser().resolve()
    new_path = args.new.expanduser().resolve()
    output = PROJECT_ROOT / "outputs" / "merged_reviews.xlsx"
    report_path = PROJECT_ROOT / "reports" / "merged_reviews_summary.md"
    temporary_paths = []
    published = []
    try:
        for destination in (output, report_path):
            if destination.resolve() in (base_path, new_path):
                raise ValueError("입력과 결과 경로가 같습니다. 원본 보호를 위해 중단합니다.")
            if destination.exists():
                raise ValueError(f"결과 파일이 이미 있습니다. 덮어쓰지 않습니다: {destination}")
        base_records, base_info = read_source(base_path, args.base_sheet)
        new_records, new_info = read_source(new_path, args.new_sheet)
        stats, distributions = diagnostics(base_records, new_records)
        report = build_report(base_info, new_info, stats, distributions)
        # Complete both temporary files before publishing either final file.
        for destination in (output, report_path):
            destination.parent.mkdir(parents=True, exist_ok=True)
            with tempfile.NamedTemporaryFile(mode="w+b", dir=destination.parent,
                                             prefix=destination.stem + ".",
                                             suffix=".tmp" + destination.suffix, delete=False) as stream:
                temporary_paths.append(Path(stream.name))
                if destination == output:
                    write_workbook(stream, base_records, new_records)
                else:
                    stream.write(report.encode("utf-8"))
                stream.flush()
                os.fsync(stream.fileno())
        for temporary_path, destination in zip(temporary_paths, (output, report_path)):
            publish_no_replace(temporary_path, destination)
            published.append(destination)
        print(report)
        print(f"병합 Excel: {output}\n요약 보고서: {report_path}")
    except PermissionError:
        print("오류: 파일 접근/저장 권한을 확인하고 열려 있는 결과 파일을 닫아주세요.", file=sys.stderr)
        return 1
    except (BadZipFile, InvalidFileException):
        print("오류: 입력 파일이 정상적인 Excel(.xlsx) 파일인지 확인해주세요.", file=sys.stderr)
        return 1
    except Exception as exc:
        print(f"오류: 리뷰 병합을 완료하지 못했습니다. {exc}", file=sys.stderr)
        return 1
    finally:
        if len(published) == 1:
            print(f"알림: Excel 저장은 완료되었으나 보고서 게시에 실패했습니다. 완성된 파일은 유지합니다: {published[0]}", file=sys.stderr)
        for temporary_path in temporary_paths:
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError as exc:
                print(f"경고: 이번 실행의 임시 파일 정리 실패: {temporary_path} ({exc})", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
