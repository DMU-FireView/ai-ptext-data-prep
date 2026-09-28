"""Build V2 without rewriting V1, source reviews, or human labels."""

import argparse
from collections import Counter, defaultdict
import json
import os
from pathlib import Path
import random
import re
import tempfile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from audit_ptext_training_dataset import read_dataset, exact_stats, conflicts, lengths, document_phrases, leakage_stats, table
from build_ptext_training_dataset import (
    PROJECT_ROOT, COLUMNS, TYPES, Groups, assign_splits, fingerprint, read_inputs,
    union_equal, union_similar, summarize, validate, is_blank,
)
from generate_ptext_synthetic_v2 import generate_v2, starting_key


def deduplicate_normal(rows, seed=42):
    """Keep human-reviewed content first; ties use seeded choice, never edit text."""
    rng = random.Random(seed)
    groups = defaultdict(list)
    for row in rows:
        if row["label"] == "NORMAL" and row["label_origin"] != "synthetic":
            groups[row["content"]].append(row)
    keep, removals = [], []
    rank = {"human_candidate_review": 0, "sampled_non_candidate": 1}
    for content in sorted(groups):
        members = sorted(groups[content], key=lambda row: row["source_excel_row"])
        priority = min(rank.get(row["label_origin"], 2) for row in members)
        eligible = [row for row in members if rank.get(row["label_origin"], 2) == priority]
        chosen = rng.choice(eligible)
        keep.append(dict(chosen))
        for row in members:
            if row is not chosen:
                removals.append({"removed_source_row": row["source_excel_row"],
                                 "removed_origin": row["label_origin"],
                                 "kept_source_row": chosen["source_excel_row"],
                                 "kept_origin": chosen["label_origin"]})
    keep.extend(dict(row) for row in rows if row["label"] != "NORMAL" and row["label_origin"] != "synthetic")
    return sorted(keep, key=lambda row: row["source_excel_row"]), removals


STOP_WORDS = set("있습니다 있어요 있습니다만 없어요 없습니다 입니다 이에요 예요 이네요 때문에 정도 이런 그런 있는 없는 합니다 하는 되는 위한 위한게 그리고 하지만 그래서 여기 바로 정말 아직 이번 그냥 것이 것을 것은 거예요 겁니다 같아요 같네요 같고 모두 그대로 이건 이거 한번 같은 만큼 있어도 있어서 없이 이미 매우 어떤 어떤게 필요한 모든 통해 그냥 살펴보면 보면 보면요 있겠어요 하나 하나는 한번도 거죠 됩니다 되네요 때는 다만 위해 보니 보다 하기 쉽게 하는데 모두에게".split())


def content_word_frequencies(rows):
    """Heuristic Korean token/josa cleanup, not a morphological analyzer."""
    counts = Counter()
    for row in rows:
        tokens = set()
        for word in re.findall(r"[가-힣A-Za-z]{2,}", row["content"]):
            if word in STOP_WORDS:
                continue
            word = re.sub(r"(?:에서는|에서도|으로는|이라는|라는|이라서|에서|에게|까지|처럼|만큼|으로|에는|에도|은|는|이|가|을|를|의|와|과|도|만)$", "", word)
            if len(word) >= 2 and word not in STOP_WORDS:
                tokens.add(word)
        counts.update(tokens)
    return counts


def quality_audit(rows):
    synthetic = [row for row in rows if row["label_origin"] == "synthetic"]
    normal = [row for row in rows if row["label"] == "NORMAL"]
    repeated = conflicts(rows)
    leakage = leakage_stats(rows)
    exact_splits = defaultdict(set)
    for row in rows:
        exact_splits[row["content"]].add(row["split"])
    cross_exact = sum(len(splits) > 1 for splits in exact_splits.values())
    starts, peak = {}, 0.0
    for kind in TYPES:
        members = [row for row in synthetic if row["synthetic_type"] == kind]
        entries = {}
        for size in (3, 4):
            counts = Counter(starting_key(row["content"], size) for row in members)
            ratio = max(counts.values(), default=0) / len(members) if members else 0
            entries[str(size)] = {"peak_ratio": ratio, "top": counts.most_common(10)}
            peak = max(peak, ratio)
        starts[kind] = entries
    rates = {word: sum(word in row["content"] for row in synthetic) / len(synthetic) if synthetic else 0
             for word in ("선택", "구성")}
    sl, nl = lengths(synthetic), lengths(normal)
    gap = abs(sl["mean"] - nl["mean"]) / nl["mean"] if nl["mean"] else 0
    words = content_word_frequencies(synthetic)
    warnings = [(word, count, count / len(synthetic)) for word, count in words.most_common()
                if synthetic and count / len(synthetic) >= .30]
    by_source = {row["source_excel_row"]: row for row in rows if row["label_origin"] != "synthetic"}
    ratios = [len(row["content"]) / len(by_source[row["parent_source_excel_row"]]["content"])
              for row in synthetic if row["parent_source_excel_row"] in by_source]
    mandatory = {
        "label_conflict_zero": not repeated,
        "group_leakage_zero": leakage["group_id"] == 0,
        "parent_leakage_zero": leakage["parent_review_id"] == 0 and leakage["parent_source_row"] == 0,
        "family_leakage_zero": leakage["synthetic_family_id"] == 0,
        "cross_split_exact_zero": cross_exact == 0,
        "uncertain_zero": all(row["label"] != "UNCERTAIN" for row in rows),
        "existing_validator": leakage["existing_validator"] == "PASS",
    }
    targets = {"선택_below_30pct": rates["선택"] < .30, "구성_below_30pct": rates["구성"] < .30,
               "prefix_below_5pct": peak < .05, "mean_length_gap_within_20pct": gap <= .20}
    family_counts = Counter(row["synthetic_family_id"] for row in synthetic if row["synthetic_type"] == "COPY_VARIATION")
    return {
        "mandatory": mandatory, "targets": targets, "label_conflict_groups": len(repeated),
        "exact": {"all": exact_stats(rows), "NORMAL": exact_stats(normal), "synthetic": exact_stats(synthetic)},
        "leakage": leakage, "cross_split_exact_content": cross_exact,
        "lengths": {"synthetic": sl, "real_NORMAL": nl,
                    **{kind: lengths([r for r in synthetic if r["synthetic_type"] == kind]) for kind in TYPES}},
        "mean_length_gap": gap, "monitored_word_rates": rates, "prefix_peak_ratio": peak,
        "prefixes": starts, "frequent_phrases": document_phrases(synthetic),
        "content_word_top30": words.most_common(30), "content_word_warnings": warnings,
        "parent_length_ratio_min": min(ratios, default=0), "parent_length_ratio_max": max(ratios, default=0),
        "parent_length_in_range_count": sum(.7 <= ratio <= 1.3 for ratio in ratios),
        "copy_family_max": max(family_counts.values(), default=0),
    }


def build_v2(v1, reviews, labels, synthetic_count=1250, seed=42):
    old = read_dataset(v1)
    real, removed = deduplicate_normal(old, seed)
    raw, human_labels = read_inputs(reviews, labels)
    original = {row["source_excel_row"]: row for row in raw}
    for row in real:
        source = original.get(row["source_excel_row"])
        if source is None or any(source[field] != row[field] for field in ("platform", "product_id", "review_id", "content")):
            raise ValueError("v1 실제 행과 원본 리뷰 대조 실패")
        if human_labels.get(row["source_excel_row"]) == "UNCERTAIN":
            raise ValueError("UNCERTAIN source가 v1 실제 행에 포함되어 있습니다.")
    generated = generate_v2(real, synthetic_count, seed)
    old_synthetic_texts = {r["content"] for r in old if r["label_origin"] == "synthetic"}
    if any(r["content"] in old_synthetic_texts for r in generated):
        raise ValueError("v1 synthetic 본문과 동일한 v2 본문이 생성되었습니다.")
    valid_raw = [r for r in raw if isinstance(r["content"], str) and not is_blank(r["content"])]
    index = {row["source_excel_row"]: i for i, row in enumerate(valid_raw)}
    graph = Groups(len(valid_raw) + len(generated))
    union_equal(graph, ((i, r["review_id"]) for i, r in enumerate(valid_raw) if not is_blank(r["review_id"])))
    for i, row in enumerate(generated, len(valid_raw)):
        graph.union(i, index[row["parent_source_excel_row"]])
    edges = union_similar(graph, [r["content"] for r in valid_raw + generated], .85)
    rows = real + generated
    for row in rows:
        parent_number = row["parent_source_excel_row"] if row["label_origin"] == "synthetic" else row["source_excel_row"]
        row["group_id"] = f"v2-group-{valid_raw[graph.find(index[parent_number])]['source_excel_row']:06d}"
        row["seed"] = seed
    assign_splits(rows, seed)
    for row in rows:
        if row["split"] != "train" and row["label"] == "NORMAL":
            row["evaluation_set"] = "real_NORMAL_human" if row["label_origin"] == "human_candidate_review" else "real_NORMAL_weak"
    audit = quality_audit(rows)
    metadata = {
        "seed": seed, "v1_rows": len(old), "v1_synthetic_discarded": sum(r["label_origin"] == "synthetic" for r in old),
        "removed_exact_normal": len(removed), "removed_by_origin": dict(Counter(r["removed_origin"] for r in removed)),
        "retention_rule": "human_candidate_review first; seeded tie choice; raw content exact match only",
        "removed_records": removed, "near_threshold": .85, "near_edges": edges,
        "uncertain_excluded_source": sum(label == "UNCERTAIN" for label in human_labels.values()),
        "new_synthetic_count": len(generated), "all_parent_count": len({r['parent_source_excel_row'] for r in generated}),
        "copy_family_limit": 2,
        "platform_label_split": dict(Counter(" | ".join((row["platform"], row["label"], row["split"])) for row in rows)),
    }
    return rows, summarize(rows), audit, metadata


def report_text(summary, audit, metadata, hashes, detailed=False):
    lines = ["# P_text training dataset v2 " + ("audit" if detailed else "summary"), "",
             "모델 입력은 content만 사용합니다. provenance/ID/platform/synthetic_type/split은 모델 입력이 아닙니다.",
             "v1·원본 리뷰·사람 라벨은 수정하지 않았습니다. 이 단계는 학습 실행이나 실제 positive 성능 검증이 아닙니다.", "",
             f"- seed: {metadata['seed']}", f"- NORMAL exact 추가 중복 제거: {metadata['removed_exact_normal']}행",
             "- 보존 기준: 원문이 정확히 같을 때 human_candidate_review 우선, 동순위는 seed 42로 1행 결정. 별개 사용자의 동일한 표현일 수 있으며 위험 판정이나 원본 삭제를 뜻하지 않습니다.",
             f"- v1 synthetic {metadata['v1_synthetic_discarded']}행은 사용하지 않고 새로 생성했습니다.", ""]
    for field, values in summary.items():
        lines += [f"## {field}", ""]
        lines += table(["항목", "수"], values.items() if isinstance(values, dict) else [(field, values)])
    lines += ["## platform × label × split", ""]
    triples = metadata["platform_label_split"]
    lines += table(["platform / label / split", "수"], triples.items())
    lines += ["## Quality gate", ""]
    for category in ("mandatory", "targets"):
        lines += table([category, "결과"], [(name, "PASS" if passed else "FAIL") for name, passed in audit[category].items()])
    lines += [f"- 선택 등장률: {audit['monitored_word_rates']['선택']:.2%}",
              f"- 구성 등장률: {audit['monitored_word_rates']['구성']:.2%}",
              f"- 유형 내 최빈 시작 구문 비율 (첫 3/4어절 중 최대): {audit['prefix_peak_ratio']:.2%}",
              f"- 실제 NORMAL 대비 synthetic 평균 길이 차이: {audit['mean_length_gap']:.2%}",
              f"- parent 길이 대비 70~130% 범위: {audit['parent_length_in_range_count']}/{metadata['new_synthetic_count']}행",
              f"- parent 길이 비율 최소/최대: {audit['parent_length_ratio_min']:.4f}/{audit['parent_length_ratio_max']:.4f}",
              f"- COPY_VARIATION family 최대 크기: {audit['copy_family_max']}", "",
              "## 길이 분포", ""]
    lines += table(["유형", "수", "평균", "중앙값", "최소", "최대"],
                   [(name, *(item[key] for key in ("count", "mean", "median", "min", "max"))) for name, item in audit["lengths"].items()])
    lines += ["## 충돌 / 중복 / 누수", "", f"- label conflict 그룹: {audit['label_conflict_groups']}",
              f"- cross-split exact-content: {audit['cross_split_exact_content']}", ""]
    lines += table(["범위", "exact 그룹", "참여 행", "추가 행"],
                   [(scope, stats["groups"], stats["participating_rows"], stats["extra_rows"]) for scope, stats in audit["exact"].items()])
    lines += table(["leakage 검사", "결과"], audit["leakage"].items())
    lines += ["## 비기능어 편중 진단", "",
              "문서별 포함 횟수(한 행당 1회)를 셉니다. 정규식 토큰화·간단한 조사 제거·기능어 목록에 의한 근사이며 한국어 형태소 분석이 아닙니다. 미탐지/오탐이 가능합니다.", ""]
    if audit["content_word_warnings"]:
        lines += ["30% 이상 비기능어 후보 경고:", ""]
        lines += table(["단어", "포함 행", "비율"], [(w, n, f"{rate:.2%}") for w, n, rate in audit["content_word_warnings"]])
    else:
        lines += ["이 근사 진단에서 30% 이상인 비기능어 후보는 없습니다.", ""]
    lines += table(["상위 단어", "포함 행"], audit["content_word_top30"])
    lines += ["## 반복 표현 상위 항목", ""]
    lines += table(["구문", "문서 빈도"], audit["frequent_phrases"])
    lines += ["## 시작 구문", "",
              "문장부호를 제거하고 첫 3/4어절을 비교해 표면적으로 가까운 시작을 함께 검사합니다. 의미적으로 유사한 모든 시작을 탐지한다는 뜻은 아닙니다.", ""]
    for kind, sizes in audit["prefixes"].items():
        lines += [f"### {kind}", ""]
        lines += table(["어절 수", "최빈 비율"], [(size, f"{entry['peak_ratio']:.2%}") for size, entry in sizes.items()])
        if detailed:
            lines += table(["시작 3어절", "수"], sizes["3"]["top"])
    lines += ["## 평가 및 한계", "",
              "- Real NORMAL evaluation: validation/test의 human_candidate_review NORMAL과 sampled_non_candidate NORMAL을 따로 집계합니다. false positive rate 확인 목적이며 후자는 사람이 확인하지 않은 약한 라벨이라 참고용입니다.",
              "- Synthetic challenge evaluation: validation/test의 synthetic SUSPICIOUS만 별도로 집계합니다. 실제 positive recall의 대체 지표가 아닙니다.",
              "- 실제 SUSPICIOUS는 1건뿐이므로 실제 positive recall을 신뢰성 있게 평가할 수 없습니다. 데이터셋 생성 중 모델 학습/평가를 실행하지 않았습니다.",
              "- group 단위 70/15/15 목표이며 실제 비율은 연결 그룹 크기에 따라 달라집니다. 모든 원본과 v2 synthetic의 공백 정리 exact 및 15자 이상 char_wb 3~5gram cosine≥0.85 연결요소를 사용합니다. 부모 ID와 원본 행, family를 함께 묶습니다.",
              "- v1 synthetic 문장은 사용하지 않습니다. v2는 실제 seed에서 상품군을 파악하고 유형별 의미 절, 문체, 순서, 권유 방식과 길이를 다르게 조합합니다. 생성 주장은 제품에 대한 검증된 사실이 아닙니다.",
              "- parent는 유지된 실제 리뷰 중 60~600자에서 플랫폼 비례로 뽑습니다. 홍보·COPY 유형은 일반적인 칭찬을 장황하게 늘이지 않도록 300자 이하의 parent만 사용합니다. 짧거나 긴 실제 학습 행을 삭제하지는 않습니다. 길이 제약은 의미 있는 별도 설명/주장 절로 맞추며 무의미한 패딩이나 문자열 잘라내기는 하지 않습니다.",
              "- 정량 목표 통과는 문장 자연스러움·유형 적합성·임상/제품 사실성 검증이 아닙니다. 조합형 생성의 문체 편향이 완전히 사라졌다고 보장할 수 없습니다. 실제 리뷰가 같은 표현을 쓴다는 이유만으로 의심 라벨을 붙이지 마세요.", "",
              "## 보존 파일 SHA256", ""]
    lines += [f"- {path}: `{value}`" for path, value in hashes.items()]
    if detailed:
        lines += ["", "## v2에서 제외된 NORMAL 중복 행 기록", "",
                  "행 번호는 원본 merged_reviews.xlsx의 source_excel_row입니다. 원본은 보존되어 있습니다.", ""]
        lines += table(["제외 원본 행", "제외 origin", "보존 원본 행", "보존 origin"],
                       [(r["removed_source_row"], r["removed_origin"], r["kept_source_row"], r["kept_origin"]) for r in metadata["removed_records"]])
    return "\n".join(lines) + "\n"


def save_bundle(rows, summary, audit, metadata, hashes, output, report, audit_path):
    if not all(audit["mandatory"].values()):
        raise ValueError("필수 quality gate 실패: " + json.dumps(audit["mandatory"], ensure_ascii=False))
    paths = (output, report, audit_path)
    if len(set(paths)) != 3 or any(path.exists() for path in paths):
        raise ValueError("출력 경로가 겹치거나 기존 결과가 있습니다. 덮어쓰지 않습니다.")
    temporary = []
    try:
        for path in paths:
            path.parent.mkdir(parents=True, exist_ok=True)
            with tempfile.NamedTemporaryFile(mode="w+b", dir=path.parent, prefix=path.stem + ".", suffix=".tmp", delete=False) as stream:
                temporary.append(Path(stream.name))
                if path == output:
                    workbook = Workbook(write_only=True)
                    try:
                        sheet = workbook.create_sheet("dataset")
                        sheet.append(list(COLUMNS))
                        for row in rows:
                            cells = []
                            for field in COLUMNS:
                                cell = WriteOnlyCell(sheet, value=row.get(field))
                                if isinstance(row.get(field), str):
                                    cell.data_type = "s"
                                cells.append(cell)
                            sheet.append(cells)
                        sheet = workbook.create_sheet("summary")
                        sheet.append(["section", "item", "value"])
                        for section, values in summary.items():
                            for item, value in (values.items() if isinstance(values, dict) else [(section, values)]):
                                sheet.append([section, str(item), value])
                        for category in ("mandatory", "targets"):
                            for item, passed in audit[category].items():
                                sheet.append([category, item, "PASS" if passed else "FAIL"])
                        workbook.save(stream)
                    finally:
                        workbook.close()
                else:
                    stream.write(report_text(summary, audit, metadata, hashes, path == audit_path).encode("utf-8"))
                stream.flush()
                os.fsync(stream.fileno())
        for temporary_path, path in zip(temporary, paths):
            if os.name == "nt":
                temporary_path.rename(path)
            else:
                os.link(temporary_path, path)
    finally:
        for path in temporary:
            path.unlink(missing_ok=True)


def main():
    parser = argparse.ArgumentParser(description="v1 보존 + NORMAL exact 정리 + 새 synthetic v2")
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--synthetic-count", type=int, default=1250)
    args = parser.parse_args()
    v1 = PROJECT_ROOT / "outputs/ptext_training_dataset_v1.xlsx"
    reviews = PROJECT_ROOT / "outputs/merged_reviews.xlsx"
    labels = PROJECT_ROOT / "outputs/merged_ptext_candidates_human_review_labeled_v1.xlsx"
    protected = [v1, reviews, labels, PROJECT_ROOT / "reports/ptext_training_dataset_v1_summary.md",
                 PROJECT_ROOT / "reports/ptext_training_dataset_v1_audit.md"]
    destinations = [PROJECT_ROOT / "outputs/ptext_training_dataset_v2.xlsx",
                    PROJECT_ROOT / "reports/ptext_training_dataset_v2_summary.md",
                    PROJECT_ROOT / "reports/ptext_training_dataset_v2_audit.md"]
    if any(path.exists() for path in destinations):
        raise ValueError("v2 결과가 이미 있습니다. 덮어쓰지 않습니다.")
    hashes = {str(path.relative_to(PROJECT_ROOT)): fingerprint(path) for path in protected}
    rows, summary, audit, metadata = build_v2(v1, reviews, labels, args.synthetic_count, args.seed)
    if any(fingerprint(path) != hashes[str(path.relative_to(PROJECT_ROOT))] for path in protected):
        raise ValueError("보존 대상 파일 해시 변경")
    print(json.dumps({"summary": summary, "removed": metadata["removed_exact_normal"],
                      "mandatory": audit["mandatory"], "targets": audit["targets"],
                      "word_rates": audit["monitored_word_rates"], "prefix_peak": audit["prefix_peak_ratio"],
                      "lengths": audit["lengths"], "content_word_warnings": audit["content_word_warnings"]}, ensure_ascii=False, indent=2))
    save_bundle(rows, summary, audit, metadata, hashes, *destinations)
    if any(fingerprint(path) != hashes[str(path.relative_to(PROJECT_ROOT))] for path in protected):
        raise ValueError("저장 후 보존 대상 파일 해시 변경")
    print("v2 저장 완료. v1/원본/사람 라벨 SHA256 모두 유지. 목표 실패 항목은 위 targets 및 보고서에 표시합니다.")


if __name__ == "__main__":
    main()
