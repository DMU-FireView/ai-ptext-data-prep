"""Inspect raw Reviews without changing source values or the input workbook."""

import argparse
from pathlib import Path
import sys

import pandas as pd
from openpyxl import Workbook, load_workbook


COLUMNS = (
    "platform", "product_id", "review_id", "content", "rating", "author",
    "written_at", "option", "images", "helpful_count", "collected_at",
)
PROJECT_ROOT = Path(__file__).resolve().parent.parent


def is_blank(value):
    """Use stripping only for the predicate; never replace stored values."""
    return value is None or (isinstance(value, str) and not value.strip()) or bool(pd.isna(value))


def read_reviews(path):
    if not path.is_file():
        raise ValueError(f"입력 파일이 없습니다: {path}")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if "Reviews" not in workbook.sheetnames:
            raise ValueError("입력 Excel에 'Reviews' 시트가 없습니다.")
        rows = list(workbook["Reviews"].iter_rows(values_only=True))
    finally:
        workbook.close()

    header_index = None
    positions = None
    header_count = 0
    for index, row in enumerate(rows):
        # Only header labels are stripped to recognize the table structure.
        labels = [value.strip() if isinstance(value, str) else value for value in row]
        if all(labels.count(column) == 1 for column in COLUMNS):
            header_index = index
            positions = [labels.index(column) for column in COLUMNS]
            header_count += 1
    if header_index is None:
        raise ValueError("기대 header를 찾지 못했습니다. 필요한 컬럼: " + ", ".join(COLUMNS))

    records = []
    source_rows = []
    skipped = 0
    for number, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        record = [row[position] if position < len(row) else None for position in positions]
        if all(is_blank(value) for value in record):
            skipped += 1
            continue
        records.append(record)
        source_rows.append(number)
    # Object dtype avoids coercing identifiers or other original cell values.
    frame = pd.DataFrame(records, columns=COLUMNS, dtype=object)
    frame.index = pd.Index(source_rows, name="source_excel_row")
    return frame, header_index + 1, header_count, skipped


def inspect(frame):
    blanks = frame.apply(lambda column: column.map(is_blank)).astype(bool)
    product_key = ["platform", "product_id"]
    review_key = [*product_key, "review_id"]
    valid_product = ~blanks[product_key].any(axis=1)
    ids = frame.loc[~blanks[review_key].any(axis=1)]
    contents = frame.loc[~blanks["content"]]
    products = frame.loc[valid_product]
    identified_contents = frame.loc[valid_product & ~blanks["content"]]

    id_duplicates = ids.loc[ids.duplicated(subset=review_key, keep=False)]
    id_sizes = ids.groupby(review_key, sort=False).size()
    content_duplicates = contents.loc[contents["content"].duplicated(keep=False)]
    content_sizes = contents.groupby("content", sort=False).size()
    within_sizes = identified_contents.groupby([*product_key, "content"], sort=False).size()
    product_counts = identified_contents.drop_duplicates(
        subset=[*product_key, "content"]
    ).groupby("content", sort=False).size()
    cross_contents = product_counts.index[product_counts > 1]
    cross_duplicates = identified_contents.loc[identified_contents["content"].isin(cross_contents)]

    metrics = [
        ("전체 리뷰 수", len(frame)),
        ("고유 상품 수 ((platform, product_id), 빈값 제외)", len(products.drop_duplicates(subset=product_key))),
        ("review_id 빈값 수", blanks["review_id"].sum()),
        ("복합 리뷰키 중복 수 (최초 1행 제외 추가 행)", ids.duplicated(subset=review_key).sum()),
        ("복합 리뷰키 중복 그룹 수", (id_sizes > 1).sum()),
        ("복합 리뷰키 중복 참여 전체 행 수", len(id_duplicates)),
        ("content 빈값 수", blanks["content"].sum()),
        ("완전히 동일한 content 중복 그룹 수", (content_sizes > 1).sum()),
        ("동일 content로 인해 추가 중복된 리뷰 행 수", contents["content"].duplicated().sum()),
        ("동일 content 중복 참여 전체 행 수", len(content_duplicates)),
        ("동일 상품 (platform, product_id) 안의 동일 content 반복 그룹 수", (within_sizes > 1).sum()),
        ("서로 다른 상품 (platform, product_id) 사이의 동일 content 반복 그룹 수", len(cross_contents)),
    ]
    details = {
        "duplicate_review_ids": id_duplicates,
        "duplicate_contents": content_duplicates,
        # Excel worksheet names are limited to 31 characters.
        "cross_product_duplicate_content": cross_duplicates,
    }
    return blanks, metrics, details


def markdown_value(value):
    # Escaping affects report presentation only, never the dataframe/export.
    return str(value).replace("&", "&amp;").replace("<", "&lt;").replace(
        ">", "&gt;"
    ).replace("|", "&#124;").replace("\r", "\\r").replace("\n", "\\n")


def table(title, label, entries):
    lines = [f"## {title}", "", f"| {label} | 개수 |", "| --- | ---: |"]
    lines.extend(f"| {markdown_value(value)} | {int(count)} |" for value, count in entries)
    return lines + [""]


def identifier_diagnostics(frame):
    """Build report-only keys without assigning values back to the dataframe."""
    lines = [
        "## 식별자 타입 진단", "",
        "- Python 값 타입별 행 수는 빈값을 제외하며, 제외한 빈값 수는 별도로 표시합니다.",
        "- 동일 표현 후보는 각 ID 컬럼 전체에서 플랫폼/상품 구분 없이 진단합니다.",
        "- 진단용 비교키만 별도로 만듭니다. 문자열은 그대로, int는 십진 문자열, float는 정수값이면 정수 표현, 그 외에는 문자열 표현을 사용합니다.",
        "- 예: 123, 123.0, 문자열 '123'은 같은 진단용 표현입니다. 문자열의 공백이나 선행 0은 제거하지 않습니다.",
        "- 같은 진단용 표현에 서로 다른 Python 타입이 2개 이상 존재하면 후보 그룹으로 집계합니다. int/float/str 외 타입은 후보 검사에서 제외합니다.",
        "- 후보는 타입 혼재 경고이며 실제 복합키 중복이나 같은 상품/리뷰임을 확정하지 않습니다. int와 float는 기존 비교에서 같게 처리될 수도 있습니다.",
        "- 원본 ID, content, 기존 중복 판정은 변경하지 않으며 자동 병합하지 않습니다. 후보 ID 목록은 표시하지 않습니다.", "",
    ]
    for column in ("product_id", "review_id"):
        type_counts = {}
        representations = {}
        blank_count = 0
        for value in frame[column]:
            if is_blank(value):
                blank_count += 1
                continue
            value_type = type(value)
            type_name = value_type.__name__
            type_counts[type_name] = type_counts.get(type_name, 0) + 1
            if value_type is str:
                comparison_key = value
            elif value_type is int:
                comparison_key = str(value)
            elif value_type is float:
                comparison_key = str(int(value)) if value.is_integer() else str(value)
            else:
                continue
            types = representations.setdefault(comparison_key, {})
            types[value_type] = types.get(value_type, 0) + 1
        candidates = [types for types in representations.values() if len(types) > 1]
        lines += table(
            f"{column} Python 값 타입별 행 수", "타입",
            [*sorted(type_counts.items()), ("[빈값: 타입 집계 제외]", blank_count)],
        )
        lines += table(f"{column} 타입 혼재 가능성", "진단 항목", [
            ("서로 다른 타입의 동일 표현 ID 후보 그룹 수", len(candidates)),
            ("후보 그룹에 포함된 전체 행 수", sum(sum(types.values()) for types in candidates)),
        ])
    return lines


def build_report(path, frame, header_row, header_count, skipped, blanks, metrics):
    try:
        display_path = path.resolve().relative_to(PROJECT_ROOT).as_posix()
    except ValueError:
        display_path = str(path)
    lines = [
        "# Raw review 1차 품질검사", "",
        f"- 입력 파일: {markdown_value(display_path)}",
        "- 시트: Reviews",
        f"- 발견한 유효 header 수: {header_count}",
        f"- 선택한 마지막 header: Excel {header_row}행",
        f"- 데이터 탐색 시작: Excel {header_row + 1}행",
        f"- 공통 컬럼이 모두 빈 행 제외: {skipped}행", "",
        "## 집계 기준", "",
        "- 공통 컬럼 11개를 모두 한 번씩 가진 마지막 header 이후만 검사합니다.",
        "- 공통 컬럼이 하나라도 비어 있지 않은 행은 리뷰로 집계합니다.",
        "- None/NaN, 빈 문자열, 공백만 있는 문자열은 빈값입니다.",
        "- 원본 값은 정규화하지 않습니다. content는 공백·줄바꿈·특수문자·이모지를 포함한 정확한 값으로 비교합니다.",
        "- 리뷰 식별 중복은 (platform, product_id, review_id) 복합키 기준이며 세 값이 모두 존재하는 행만 검사합니다.",
        "- 복합 리뷰키 중복 그룹 수는 해당 키의 행 수가 2개 이상인 그룹 수입니다.",
        "- 전체 content 중복은 플랫폼/상품과 무관하게 독립적으로 계산하며 content 빈값만 제외합니다.",
        "- 추가 중복 행 수는 각 중복 그룹의 행 수에서 최초 1행을 뺀 값의 합입니다.",
        "- 고유 상품 수와 상품별 리뷰 수는 (platform, product_id) 원본 값의 조합 기준입니다.",
        "- 상품 내 반복은 (platform, product_id, content)별 그룹 수입니다.",
        "- 상품 간 반복은 서로 다른 (platform, product_id) 조합 2개 이상에 존재하는 content의 그룹 수입니다. 상품 내 반복과 겹칠 수 있습니다.",
        "- platform 또는 product_id가 빈 행은 고유 상품 수, 상품별 집계 및 상품 내/상품 간 중복 판단에서 제외합니다.",
        "- 중복 상세에는 최초 행을 포함하며 source_excel_row는 원본 Excel 행 번호입니다.",
        "- duplicate_review_ids 시트에는 복합 리뷰키 기준 중복 행을 저장합니다.",
        "- 상품 간 상세는 platform과 product_id가 모두 있는 행만 포함합니다.",
        "- 상품 간 상세 시트명은 Excel의 31자 제한에 맞춰 cross_product_duplicate_content를 사용합니다.",
        "- 수식 셀은 계산하지 않고 원본 수식 문자열로 검사합니다.",
        "- 보고서에는 리뷰 본문을 싣지 않습니다.", "",
    ]
    lines += table("주요 통계", "항목", metrics)
    for column, title in [("platform", "플랫폼별 리뷰 수"), ("product_id", "상품별 리뷰 수"), ("rating", "rating 분포")]:
        if column == "product_id":
            product_key = ["platform", "product_id"]
            valid_product = ~blanks[product_key].any(axis=1)
            counts = frame.loc[valid_product].groupby(product_key, sort=False).size()
            entries = list(counts.items())
            entries.append(("[상품 식별 불가: platform 또는 product_id 빈값]", (~valid_product).sum()))
            lines += table(title, "(platform, product_id)", entries)
        else:
            counts = frame.loc[~blanks[column], column].value_counts(sort=False)
            entries = list(counts.items())
            entries.append(("[빈값 합계]", blanks[column].sum()))
            lines += table(title, column, entries)
    lines += table("컬럼별 null/빈값", "컬럼", blanks.sum().items())
    lines += identifier_diagnostics(frame)
    return "\n".join(lines)


def write_duplicates(path, details):
    workbook = Workbook()
    for index, (name, frame) in enumerate(details.items()):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = name
        sheet.append(["source_excel_row", *COLUMNS])
        for row_number, values in enumerate(frame.itertuples(index=True, name=None), start=2):
            for column_number, value in enumerate(values, start=1):
                if value is None or (not isinstance(value, str) and pd.isna(value)):
                    value = None
                cell = sheet.cell(row=row_number, column=column_number, value=value)
                if isinstance(value, str):
                    # Keep strings starting with '=' as literal source text.
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    try:
        workbook.save(path)
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description="Reviews 시트의 마지막 공통 header 이후 raw review 품질을 검사합니다.")
    parser.add_argument("excel_path", type=Path, help="입력 Excel 경로 (원본은 읽기 전용)")
    parser.add_argument("--report", type=Path, default=Path("reports/review_quality_summary.md"),
                        help="요약 보고서 경로 (상대경로는 PROJECT_ROOT 기준, 절대경로 허용)")
    parser.add_argument("--duplicates", type=Path, default=Path("outputs/review_duplicates.xlsx"),
                        help="중복 상세 Excel 경로 (상대경로는 PROJECT_ROOT 기준, 절대경로 허용)")
    args = parser.parse_args()
    source = args.excel_path.expanduser().resolve()
    report_path = (PROJECT_ROOT / args.report.expanduser()).resolve()
    duplicates_path = (PROJECT_ROOT / args.duplicates.expanduser()).resolve()
    try:
        if source in (report_path.resolve(), duplicates_path.resolve()):
            raise ValueError("입력 파일과 결과 저장 경로가 같습니다. 원본 보호를 위해 중단합니다.")
        frame, header_row, header_count, skipped = read_reviews(source)
        blanks, metrics, details = inspect(frame)
        report = build_report(source, frame, header_row, header_count, skipped, blanks, metrics)
        report_path.parent.mkdir(parents=True, exist_ok=True)
        duplicates_path.parent.mkdir(parents=True, exist_ok=True)
        write_duplicates(duplicates_path, details)
        report_path.write_text(report, encoding="utf-8")
        print(report)
        print(f"\n요약 보고서: {report_path}\n중복 상세: {duplicates_path}")
    except PermissionError as exc:
        print(f"오류: 파일 접근 권한을 확인하고 열려 있는 결과 Excel을 닫아주세요. {exc}", file=sys.stderr)
        return 1
    except Exception as exc:
        print(f"오류: 품질검사를 완료하지 못했습니다. {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
