"""Extract explainable human-review candidates, never final training labels."""

import argparse
from collections import Counter, defaultdict
import math
from pathlib import Path
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

from inspect_reviews import PROJECT_ROOT, is_blank, markdown_value, read_reviews


REASONS = ("REPETITION", "PRODUCT_COPY", "STRUCTURED_INFO", "PROMOTIONAL_CTA", "HARD_NEGATIVE")
REPETITION_MIN_LENGTH = 15
FIELDS = ("platform", "product_id", "review_id", "content", "rating")
SPEC = r"용량|규격|구성|소재|사이즈|색상|제조국|사용\s*방법|성분|원산지|중량|재질"
SPEC_WORDS = re.compile(SPEC)
SPEC_ITEMS = re.compile(rf"(?:^|[\n,;/])\s*(?:[-•]\s*)?({SPEC})\s*[:：=]\s*\S+", re.M)
HEADINGS = re.compile(r"(?:^|[\n/;]|\s)(장점|단점|총평|구매\s*이유|배송|사용감|재구매)\s*[:：]\s*\S+")
NUMBERED = re.compile(r"^\s*(\d{1,2})[.)]\s+\S+", re.M)
MARKERS = re.compile(r"^\s*[✅✔▶]\ufe0f?\s*\S+", re.M)
CTA_RULES = (
    ("꼭 사세요", r"꼭\s*사세요"),
    ("강추합니다", r"강추\s*합니다"),
    ("무조건 추천", r"무조건\s*추천(?:합니다|해요|드려요)?"),
    ("구매하세요", r"구매\s*하세요"),
    ("쟁이세요", r"쟁이세요"),
    ("꼭 써보세요", r"꼭\s*써\s*보세요"),
    ("추천드립니다", r"추천\s*드립니다"),
)
ACQUISITION = (
    ("추천받아 구매", r"추천\s*받(?:아서|아|고)[^.\n!?]{0,20}(?:구매|샀|샀어요)"),
    ("친구 추천으로 구매", r"친구가?\s*추천(?:해서|해\s*줘서)[^.\n!?]{0,20}(?:구매|샀)"),
    ("SNS/영상 보고 구매", r"(?:인스타|유튜브)(?:그램)?(?:에서)?[^.\n!?]{0,20}보고[^.\n!?]{0,20}(?:구매|샀)"),
    ("광고 보고 구매", r"광고\s*(?:를\s*)?보고[^.\n!?]{0,20}(?:구매|샀)"),
    ("SNS/추천템/광고 접한 후 구매", r"(?:(?i:SNS)|인스타그램|인스타|유튜브|추천템|광고)[^.\n!?]{0,20}(?:보고|접하고|소개\s*(?:되어|돼서|받아))[^.\n!?]{0,20}(?:구매|샀)"),
    ("소개받아 구매", r"소개\s*(?:되어|돼서|받아)[^.\n!?]{0,20}(?:구매|샀)"),
)
EXPERIENCE = (
    ("사용 후 관찰", r"(?:써\s*보니|사용해\s*보니|사용했는데|써봤는데|발라\s*보니|먹어\s*보니|입어\s*보니|신어\s*보니)[^.\n!?]{2,}"),
    ("직접 사용 경험", r"(?:사용해\s*봤어요|써\s*봤어요|사용해\s*보니|써\s*보니)"),
    ("만족 경험", r"(?<![가-힣])만족(?:해요|합니다|스러워요)"),
    ("마음에 든 경험", r"(?:마음|맘)에\s*(?:들어요|들었어요|듭니다)"),
    ("디자인 평가", r"디자인(?:이|은|도)?\s*(?:좋(?:다|아요|았습니다|았어요|네요)|예쁘(?:다|네요|고)|예뻐요|예쁩니다|깔끔(?:하다|해요|합니다))"),
    ("재질 평가", r"재질(?:이|은|도)?\s*(?:좋(?:다|아요|았습니다|았어요|네요)|탄탄(?:하다|해요|합니다))"),
    ("사용감 평가", r"사용감(?:이|은|도)?\s*(?:좋(?:다|아요|았습니다|았어요|네요)|편(?:하다|해요|합니다))"),
    ("성능 평가", r"성능(?:이|은|도)?\s*좋(?:다|아요|았습니다|았어요|네요)"),
    ("기간을 둔 사용", r"(?:\d+|한|두|세)\s*(?:일|주|개월|달)(?:간|째|\s*동안)?\s*(?:사용|썼|써|먹|입|신)"),
    ("구체적 배송 경험", r"(?:배송|도착|포장)[^.\n!?]{0,20}(?:하루|이틀|사흘|늦었|찢어|터져|파손|왔어요|왔는데|빨랐|빠르네요|빠르게\s*왔|빠른\s*편이었)"),
    ("개인 상황", r"(?:저는|제가|우리\s*아이|제\s*(?:피부|발|체형))[^.\n!?]{0,30}(?:민감|건성|지성|알레르기|발볼|출근|육아)"),
    ("장단점 경험", r"(?:아쉬웠|아쉬워요|불편했|불편해요|편했|편해요|따가웠|따가워요|가려웠|가려워요|잘\s*맞았|향이\s*강했|보습이\s*좋았)"),
)


def repetition_features(frame):
    raw_groups = defaultdict(list)
    normalized_groups = defaultdict(list)
    features = defaultdict(list)
    for row, content in frame["content"].items():
        if isinstance(content, str) and not is_blank(content):
            normalized = " ".join(content.split())
            if len(normalized) < REPETITION_MIN_LENGTH:
                continue
            raw_groups[content].append(row)
            normalized_groups[normalized].append(row)
    for rows in raw_groups.values():
        if len(rows) > 1:
            for row in rows:
                features[row].append(f"원문 동일 {len(rows)}행")
    for rows in normalized_groups.values():
        if len(rows) > 1 and len({frame.at[row, "content"] for row in rows}) > 1:
            for row in rows:
                features[row].append(f"공백·줄바꿈 정리 후 동일 {len(rows)}행")
    return features


def equal_source_value(left, right):
    if is_blank(left) or is_blank(right):
        return is_blank(left) and is_blank(right)
    # No identifier string conversion; use the existing raw-value comparison rule.
    return bool(left == right)


def near_duplicate_features(path, frame):
    features = defaultdict(list)
    status = {"used": False, "accepted": 0, "skipped": 0, "message": "결과 파일 없음: raw 반복 검사만 사용"}
    if not path.is_file():
        return features, status
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        if "near_duplicate_pairs" not in workbook.sheetnames:
            raise ValueError("near_duplicate_pairs 시트 없음")
        rows = workbook["near_duplicate_pairs"].iter_rows(values_only=True)
        header = list(next(rows, ()))
        required = [f"{field}_{side}" for side in ("a", "b")
                    for field in ("source_excel_row", *FIELDS[:4])] + ["similarity"]
        if any(header.count(name) != 1 for name in required):
            raise ValueError("near-duplicate 필수 컬럼 누락 또는 중복")
        positions = {name: header.index(name) for name in required}
        links = Counter()
        for values in rows:
            if all(is_blank(value) for value in values):
                continue
            pair = {name: values[index] if index < len(values) else None
                    for name, index in positions.items()}
            score = pair["similarity"]
            valid = (type(score) in (int, float) and math.isfinite(score) and 0 < score <= 1)
            source_rows = []
            for side in ("a", "b"):
                row = pair[f"source_excel_row_{side}"]
                if type(row) not in (int, float) or row not in frame.index:
                    valid = False
                    break
                if not all(equal_source_value(frame.at[row, field], pair[f"{field}_{side}"])
                           for field in FIELDS[:4]):
                    valid = False
                    break
                source_rows.append(row)
            if (not valid or len(set(source_rows)) != 2
                    or pair["content_a"] == pair["content_b"]):
                status["skipped"] += 1
                continue
            status["accepted"] += 1
            links.update(source_rows)
        for row, count in links.items():
            features[row].append(f"기존 near-duplicate 결과 {count}pair 참여 (원본 대조 완료)")
        status["used"] = True
        status["message"] = "결과 읽음: 현재 원본과 양쪽 행이 일치하는 pair만 참고"
    except Exception as exc:
        # Optional evidence must never prevent the independent raw-data scan.
        features.clear()
        status.update(used=False, accepted=0, skipped=0,
                      message=f"참고 생략: {type(exc).__name__} (파일 형식/접근 권한 확인 필요)")
    finally:
        if workbook is not None:
            workbook.close()
    return features, status


def pattern_features(content, brand_keywords):
    found = {}
    if not isinstance(content, str) or is_blank(content):
        return found
    # Copies for matching only; raw content remains unchanged in the export.
    text = " ".join(content.split())
    keywords = {re.sub(r"\s+", "", value) for value in SPEC_WORDS.findall(content)}
    labeled_specs = {re.sub(r"\s+", "", value) for value in SPEC_ITEMS.findall(content)}
    segments = re.split(r"[\n,;/]+", content)
    spec_segments = sum(bool(SPEC_WORDS.search(segment)) for segment in segments)
    if len(keywords) >= 3 and (len(labeled_specs) >= 2 or spec_segments >= 3):
        found["PRODUCT_COPY"] = ["서로 다른 스펙 키워드 3개 이상 + 항목형 나열", "키워드=" + ",".join(sorted(keywords))]

    headings = {re.sub(r"\s+", "", value) for value in HEADINGS.findall(content)}
    numbered = set(NUMBERED.findall(content))
    markers = len(MARKERS.findall(content))
    structures = []
    if len(headings) >= 2:
        structures.append("명시적 후기 항목=" + ",".join(sorted(headings)))
    if len(numbered) >= 3:
        structures.append("서로 다른 번호 항목 3개 이상")
    if markers >= 2:
        structures.append("체크/화살표 항목 2개 이상")
    if structures:
        found["STRUCTURED_INFO"] = structures

    ctas = []
    for name, pattern in CTA_RULES:
        # Skip quoted/reported/negated phrases; explicit exhortations only.
        for match in re.finditer(r"(?<!\w)" + pattern + r"(?![가-힣A-Za-z])", text):
            before, after = text[max(0, match.start() - 1):match.start()], text[match.end():match.end() + 20]
            if before and before in "\"'‘“":
                continue
            if re.match(r"\s*[\"'’”]", after) or re.match(r"\s*(?:라고|라는|라며|하지\s*않|안\s*해|은\s*아니|는\s*아니)", after):
                continue
            ctas.append(name)
            break
    if ctas:
        found["PROMOTIONAL_CTA"] = ctas

    exposure = [name for name, pattern in ACQUISITION if re.search(pattern, text)]
    if any(keyword in text for keyword in brand_keywords):
        exposure.append("지정 브랜드/제품명 언급")
    experiences = [name for name, pattern in EXPERIENCE if re.search(pattern, text)]
    if exposure and experiences:
        found["HARD_NEGATIVE"] = exposure + experiences
    return found


def write_candidates(path, frame, repetition, near, brand_keywords):
    counts = Counter({reason: 0 for reason in REASONS})
    total = multiple = 0
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("candidates")
    sheet.freeze_panes = "A2"
    sheet.append(["source_excel_row", *FIELDS, "candidate_reasons", "matched_features"])
    try:
        for values in frame[list(FIELDS)].itertuples(index=True, name=None):
            row = values[0]
            found = pattern_features(values[4], brand_keywords)
            repeated = repetition.get(row, []) + near.get(row, [])
            if repeated:
                found["REPETITION"] = repeated
            reasons = [reason for reason in REASONS if reason in found]
            if not reasons:
                continue
            if total >= 1_048_575:
                raise ValueError("후보가 Excel 단일 시트 행 한도를 초과했습니다.")
            evidence = "; ".join(f"{reason}: {', '.join(found[reason])}" for reason in reasons)
            cells = []
            for value in (*values, "|".join(reasons), evidence):
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"  # Preserve formula-like strings as literal text.
                cells.append(cell)
            sheet.append(cells)
            counts.update(reasons)
            total += 1
            multiple += len(reasons) > 1
        workbook.save(path)
    finally:
        workbook.close()
    return total, multiple, counts


def build_report(source, frame, header_row, header_count, skipped, total, multiple, counts, near_status, brands):
    try:
        display_path = source.relative_to(PROJECT_ROOT).as_posix()
    except ValueError:
        display_path = str(source)
    ratio = total / len(frame) * 100 if len(frame) else 0.0
    lines = [
        "# P_text 사람 검토용 패턴 후보", "",
        "이 결과는 자동 라벨이 아니라 사람 검토용 후보이며, candidate_reason은 최종 학습 라벨이 아닙니다.",
        "후보 여부로 SUSPICIOUS/NORMAL을 확정하지 않으며 원본을 삭제하거나 수정하지 않습니다. "
        "후보가 아닌 리뷰 역시 NORMAL로 확정하지 않습니다.", "",
        f"- 입력 파일: {markdown_value(display_path)}",
        f"- Reviews 마지막 유효 header: {header_row}행 (발견 {header_count}개)",
        f"- 공통 컬럼 전체 빈 행 제외: {skipped}행",
        f"- near-duplicate 결과 사용 여부: {'예' if near_status['used'] else '아니오'}",
        f"- near-duplicate 상태: {near_status['message']}",
        f"- 참고한 pair: {near_status['accepted']}개 / 원본 불일치·유효하지 않은 pair: {near_status['skipped']}개",
        f"- 지정 브랜드/제품명 키워드 수: {len(brands)}", "",
        "| 항목 | 값 |", "| --- | ---: |",
        f"| 전체 리뷰 수 | {len(frame)} |",
        f"| 후보 리뷰 수 | {total} |",
        f"| 후보 비율 | {ratio:.2f}% |",
        f"| 여러 유형 동시 후보 리뷰 수 | {multiple} |",
        f"| 후보가 아닌 리뷰 수 | {len(frame) - total} |",
    ]
    lines.extend(f"| {reason} 후보 수 | {counts[reason]} |" for reason in REASONS)
    lines += [
        "", "## 탐지 원칙", "",
        f"- REPETITION: raw 반복 탐지는 연속 공백·줄바꿈을 하나의 공백으로 정리하고 양끝 공백을 제거한 계산용 문자열 기준으로 {REPETITION_MIN_LENGTH}자 이상만 검사합니다. {REPETITION_MIN_LENGTH}자 미만의 짧고 흔한 반복 표현은 후보에서 제외합니다. 원문 또는 계산용 문자열이 2행 이상 동일할 때 후보이며 이상 여부를 확정하지 않습니다. 기존 near-duplicate 결과의 원본 대조 로직은 유지합니다.",
        "- 기존 outputs/near_duplicate_pairs.xlsx가 있으면 원본 행 번호와 platform/product_id/review_id/content를 양쪽 모두 대조합니다. review_id만으로 연결하지 않습니다. 일치한 기존 pair를 참고하며 유사도를 다시 계산하지 않습니다.",
        "- near-duplicate 파일이 없거나 읽을 수 없으면 해당 참고만 생략합니다. 파일의 생성 threshold는 재판정하지 않습니다. 사용 여부가 예여도 일치한 pair는 0개일 수 있습니다.",
        "- PRODUCT_COPY: 서로 다른 상품 스펙 키워드 3개 이상과, 2개 이상의 키:값 항목 또는 3개 이상의 구분된 스펙 항목을 함께 요구합니다. 숫자만 많다고 탐지하지 않습니다.",
        "- STRUCTURED_INFO: 장점/단점 등 명시적 항목 2종 이상, 서로 다른 줄 시작 번호 항목 3개 이상, 또는 줄 시작 체크/화살표 항목 2개 이상. 줄바꿈만으로 탐지하지 않습니다.",
        "- PROMOTIONAL_CTA: 꼭 사세요/강추합니다/무조건 추천/구매하세요/쟁이세요/꼭 써보세요/추천드립니다 등의 직접 표현. 인용·전언·일부 부정 표현은 제외합니다. 단순 추천 단어, 추천받은 경험, 재구매 의사, 추천할 만하다는 표현은 근거로 쓰지 않습니다.",
        "- HARD_NEGATIVE: 추천·SNS·광고를 통한 구매 또는 지정 브랜드/제품명 언급과, 직접 사용·만족·마음에 듦·디자인/재질/사용감/성능 평가·사용 기간·구체적 배송·개인 상황·장단점 경험 중 하나 이상이 함께 있어야 합니다. 구매 경로/브랜드 언급 단독이나 경험 표현 단독으로는 후보가 되지 않습니다.",
        "- 브랜드/제품명은 임의 추측하지 않습니다. 필요한 경우 --brand-keyword를 반복 지정하면 해당 원문 문자열의 언급을 확인합니다. 지정하지 않으면 구매 경로 표현만 사용합니다.",
        "- 한 리뷰는 여러 candidate_reasons를 가질 수 있으므로 유형별 합계는 후보 리뷰 수보다 클 수 있습니다. HARD_NEGATIVE도 최종 정상 라벨이 아닙니다.",
        "- 규칙은 보수적인 검토 시작점이며 문맥·부정·인용을 완전히 해석하지 못합니다. matched_features를 보고 사람이 판단해야 합니다.",
        "- content와 식별자 원본 값/타입을 유지합니다. 탐지용 문자열만 별도로 만들며 식별자를 강제로 문자열화하지 않습니다.",
        "- 빈 content와 비문자열 content는 본문 패턴 검사에서 제외합니다. 보고서에 리뷰 본문이나 전체 ID 목록을 싣지 않습니다.", "",
    ]
    return "\n".join(lines)


def main():
    parser = argparse.ArgumentParser(description="P_text 학습 준비를 위한 사람 검토용 패턴 후보 추출")
    parser.add_argument("excel_path", type=Path, help="읽기 전용 입력 Excel")
    parser.add_argument("--brand-keyword", action="append", default=[], help="HARD_NEGATIVE 검토용 브랜드/제품명 원문 키워드 (반복 지정 가능)")
    parser.add_argument("--near-duplicates", type=Path, default=Path("outputs/near_duplicate_pairs.xlsx"),
                        help="읽기 전용 near-duplicate Excel 경로 (상대경로는 PROJECT_ROOT 기준, 절대경로 허용)")
    parser.add_argument("--output", type=Path, default=Path("outputs/ptext_candidates.xlsx"),
                        help="후보 Excel 경로 (상대경로는 PROJECT_ROOT 기준, 절대경로 허용)")
    parser.add_argument("--report", type=Path, default=Path("reports/ptext_candidate_summary.md"),
                        help="요약 보고서 경로 (상대경로는 PROJECT_ROOT 기준, 절대경로 허용)")
    args = parser.parse_args()
    source = args.excel_path.expanduser().resolve()
    output = (PROJECT_ROOT / args.output.expanduser()).resolve()
    report_path = (PROJECT_ROOT / args.report.expanduser()).resolve()
    near_path = (PROJECT_ROOT / args.near_duplicates.expanduser()).resolve()
    try:
        if any(is_blank(keyword) for keyword in args.brand_keyword):
            raise ValueError("brand-keyword에는 비어 있지 않은 브랜드/제품명을 입력해주세요.")
        if source in (output.resolve(), report_path.resolve()):
            raise ValueError("입력과 출력 경로가 같습니다. 원본 보호를 위해 중단합니다.")
        if near_path in (output, report_path):
            raise ValueError("near-duplicate 입력과 출력 경로가 같습니다. 원본 보호를 위해 중단합니다.")
        if output == report_path:
            raise ValueError("output과 report 경로가 같습니다. 서로 다른 경로를 지정해주세요.")
        frame, header_row, header_count, skipped = read_reviews(source)
        repetition = repetition_features(frame)
        near, near_status = near_duplicate_features(near_path, frame)
        output.parent.mkdir(parents=True, exist_ok=True)
        report_path.parent.mkdir(parents=True, exist_ok=True)
        total, multiple, counts = write_candidates(output, frame, repetition, near, args.brand_keyword)
        report = build_report(source, frame, header_row, header_count, skipped,
                              total, multiple, counts, near_status, args.brand_keyword)
        report_path.write_text(report, encoding="utf-8")
        print(report)
        print(f"후보 Excel: {output}\n요약 보고서: {report_path}")
    except PermissionError as exc:
        print(f"오류: 파일 접근 권한을 확인하고 열려 있는 결과 Excel을 닫아주세요. {exc}", file=sys.stderr)
        return 1
    except Exception as exc:
        print(f"오류: 후보 추출을 완료하지 못했습니다. {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
