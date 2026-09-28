"""Copy candidate cells into a workbook for entirely manual human review."""

import argparse
import os
from pathlib import Path
import sys
import tempfile
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation


PROJECT_ROOT = Path(__file__).resolve().parent.parent
COLUMNS = (
    "source_excel_row", "platform", "product_id", "review_id", "content",
    "rating", "candidate_reasons", "matched_features",
)
LABELS = ("SUSPICIOUS", "NORMAL", "UNCERTAIN")


def prepare_workbook(source):
    if not source.is_file():
        raise ValueError(f"입력 파일이 없습니다: {source}")
    original = load_workbook(source, read_only=True, data_only=False)
    result = Workbook()
    try:
        if "candidates" not in original.sheetnames:
            raise ValueError("입력 Excel에 candidates 시트가 없습니다.")
        rows = original["candidates"].iter_rows()
        header = tuple(cell.value for cell in next(rows, ()))
        if header != COLUMNS:
            raise ValueError("candidates 시트의 헤더는 기존 8개 컬럼과 순서가 일치해야 합니다: " + ", ".join(COLUMNS))
        sheet = result.active
        sheet.title = "candidates"
        sheet.append([*COLUMNS, "human_label", "review_note"])
        count = 0
        for row_number, row in enumerate(rows, start=2):
            # No filtering, sorting, normalization, ID coercion or inferred labels.
            for column_number, source_cell in enumerate(row, start=1):
                cell = sheet.cell(row_number, column_number, value=source_cell.value)
                # Preserve literal strings (including '=...'), formulas, errors,
                # and number formats as stored by the source workbook.
                cell.data_type = source_cell.data_type
                if source_cell.number_format is not None:
                    cell.number_format = source_cell.number_format
            for column_number in range(1, 11):
                sheet.cell(row_number, column_number).alignment = Alignment(
                    vertical="top", wrap_text=column_number in (5, 7, 8, 10)
                )
            # The two manual cells are deliberately left as None.
            sheet.row_dimensions[row_number].height = 75
            count += 1
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:J{count + 1}"
        widths = (19, 16, 23, 24, 70, 10, 36, 58, 20, 45)
        for column_number, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column_number)].width = width
            cell = sheet.cell(1, column_number)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="FFF2CC" if column_number > 8 else "DCE6F1")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 30
        validation = DataValidation(
            type="list", formula1='"' + ",".join(LABELS) + '"',
            allow_blank=True, showDropDown=False,
        )
        validation.showErrorMessage = True
        validation.errorStyle = "stop"
        validation.errorTitle = "허용되지 않은 값"
        validation.error = "SUSPICIOUS, NORMAL, UNCERTAIN 중 선택하거나 빈칸으로 두세요."
        validation.showInputMessage = True
        validation.promptTitle = "사람이 직접 판단"
        validation.prompt = "candidate_reasons는 최종 라벨이 아닙니다. 원문 검토 후 선택하세요."
        sheet.add_data_validation(validation)
        validation.add(f"I2:I{max(2, count + 1)}")
        return result, count
    except Exception:
        result.close()
        raise
    finally:
        original.close()


def main():
    parser = argparse.ArgumentParser(description="후보 원문을 보존하고 빈 수동 검토 칸 2개를 추가합니다.")
    parser.add_argument("input", type=Path, nargs="?",
                        default=Path("outputs/merged_ptext_candidates_expansion_v2.xlsx"))
    parser.add_argument("--output", type=Path,
                        default=Path("outputs/merged_ptext_candidates_human_review.xlsx"))
    args = parser.parse_args()
    source = (PROJECT_ROOT / args.input.expanduser()).resolve()
    output = (PROJECT_ROOT / args.output.expanduser()).resolve()
    workbook = None
    temporary_path = None
    try:
        if source == output:
            raise ValueError("입력과 출력 경로가 같습니다. 원본 보호를 위해 중단합니다.")
        if output.exists():
            raise ValueError(f"출력 파일이 이미 있습니다. 기존 수동 검토 내용을 덮어쓰지 않습니다: {output}")
        workbook, count = prepare_workbook(source)
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(mode="w+b", dir=output.parent,
                                         prefix=output.stem + ".", suffix=".tmp.xlsx",
                                         delete=False) as stream:
            temporary_path = Path(stream.name)
            workbook.save(stream)
            stream.flush()
            os.fsync(stream.fileno())
        if os.name == "nt":
            temporary_path.rename(output)  # Atomic, and refuses an existing destination.
        else:
            os.link(temporary_path, output)  # Atomic no-overwrite publication on POSIX.
        print(f"저장 완료: {output}\n보존한 후보 행: {count}\nhuman_label/review_note: 모두 빈칸 (자동 라벨링 없음)")
    except PermissionError:
        print("오류: 파일 권한을 확인하고 Excel에서 열려 있는 파일을 닫아주세요.", file=sys.stderr)
        return 1
    except (BadZipFile, InvalidFileException):
        print("오류: 정상적인 .xlsx 파일인지 확인해주세요.", file=sys.stderr)
        return 1
    except Exception as exc:
        print(f"오류: 사람 검토용 Excel 생성 실패: {exc}", file=sys.stderr)
        return 1
    finally:
        if temporary_path is not None:
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError as exc:
                print(f"경고: 이번 실행 임시 파일 정리 실패: {temporary_path} ({exc})", file=sys.stderr)
        if workbook is not None:
            workbook.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
