"""Build a manual labeling workbook without assigning any final decisions."""

import argparse
import os
from pathlib import Path
import sys
import tempfile
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation


PROJECT_ROOT = Path(__file__).resolve().parent.parent
SOURCE_COLUMNS = (
    "source_excel_row", "platform", "product_id", "review_id", "content",
    "rating", "candidate_reasons", "matched_features",
)
MANUAL_COLUMNS = ("final_label", "confirmed_patterns", "use_for_training", "reviewer_note")
FINAL_LABELS = ("NORMAL", "SUSPICIOUS", "UNCERTAIN", "INVALID")
TRAINING_CHOICES = ("YES", "NO")
PATTERNS = ("REPETITION", "PRODUCT_COPY", "STRUCTURED_INFO", "PROMOTIONAL_CTA")
HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
MANUAL_FILL = PatternFill("solid", fgColor="FFF2CC")


def read_candidates(path):
    if not path.is_file():
        raise ValueError(f"입력 파일이 없습니다: {path}")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if "candidates" not in workbook.sheetnames:
            raise ValueError("입력 Excel에 'candidates' 시트가 없습니다.")
        rows = workbook["candidates"].iter_rows(values_only=True)
        header = list(next(rows, ()))
        missing = [name for name in SOURCE_COLUMNS if name not in header]
        if missing:
            raise ValueError("필수 컬럼 누락: " + ", ".join(missing))
        duplicate = [name for name in SOURCE_COLUMNS if header.count(name) > 1]
        if duplicate:
            raise ValueError("필수 컬럼이 중복되어 읽을 수 없습니다: " + ", ".join(duplicate))
        positions = [header.index(name) for name in SOURCE_COLUMNS]
        # No sorting, normalization, type coercion, or candidate filtering.
        return [tuple(row[index] if index < len(row) else None for index in positions)
                for row in rows]
    finally:
        workbook.close()


def set_literal(cell, value):
    cell.value = value
    if isinstance(value, str):
        # Keep source strings beginning with '=' as text, not executable formulas.
        cell.data_type = "s"


def add_dropdown(sheet, column, last_row, choices, title):
    validation = DataValidation(
        type="list", formula1='"' + ",".join(choices) + '"',
        allow_blank=True, showDropDown=False,
    )
    validation.errorTitle = "허용되지 않은 값"
    validation.error = "빈칸으로 두거나 목록의 값 중 하나를 선택해주세요."
    validation.errorStyle = "stop"
    validation.showErrorMessage = True
    validation.promptTitle = title
    validation.prompt = "사람이 검토한 뒤 직접 선택합니다. 초기값은 빈칸입니다."
    validation.showInputMessage = True
    sheet.add_data_validation(validation)
    validation.add(f"{column}2:{column}{max(2, last_row)}")


def build_labeling_sheet(workbook, records):
    sheet = workbook.active
    sheet.title = "labeling"
    sheet.append([*SOURCE_COLUMNS, *MANUAL_COLUMNS])
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 30
    widths = (18, 16, 22, 24, 65, 10, 28, 48, 19, 38, 21, 48)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = sheet.cell(1, index)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL if index <= len(SOURCE_COLUMNS) else MANUAL_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row_number, record in enumerate(records, start=2):
        for column_number, value in enumerate(record, start=1):
            set_literal(sheet.cell(row_number, column_number), value)
        # All four manual fields start empty, regardless of candidate_reasons.
        for column_number in range(1, 13):
            cell = sheet.cell(row_number, column_number)
            cell.alignment = Alignment(vertical="top", wrap_text=column_number in (5, 7, 8, 10, 12))
        sheet.row_dimensions[row_number].height = 72
    last_row = len(records) + 1
    sheet.auto_filter.ref = f"A1:L{last_row}"
    add_dropdown(sheet, "I", last_row, FINAL_LABELS, "final_label")
    add_dropdown(sheet, "K", last_row, TRAINING_CHOICES, "use_for_training")
    sheet["J1"].comment = Comment(
        "사람이 확인한 코드만 입력: " + ", ".join(PATTERNS)
        + "\n여러 코드는 |로 구분합니다. 예: REPETITION|PROMOTIONAL_CTA"
        + "\nHARD_NEGATIVE는 위험 패턴이 아니므로 입력하지 않습니다.", "Guide",
    )
    sheet["G1"].comment = Comment("검토용 힌트이며 최종 라벨 또는 확정 패턴이 아닙니다.", "Guide")
    if records:
        sheet.conditional_formatting.add(
            f"I2:I{last_row}",
            FormulaRule(formula=['LEN(TRIM($I2))=0'], fill=MANUAL_FILL),
        )


def build_guide_sheet(workbook):
    sheet = workbook.create_sheet("guide")
    rows = [
        ("구분", "코드/항목", "설명"),
        ("핵심 원칙", "수동 판단", "이 결과는 자동 라벨이 아니라 사람 검토용 후보이며, candidate_reason은 최종 학습 라벨이 아닙니다. 모든 최종 판단은 사람이 합니다."),
        ("작업 순서", "1 → 4", "원문과 근거를 읽고 → final_label 선택 → 확인한 패턴만 입력 → 학습 사용 여부와 메모 입력. guide를 먼저 읽고 labeling 시트에서 작업하세요."),
        ("final_label", "NORMAL", "일반적인 실제 사용자 경험 리뷰라고 판단. 실제 사용 경험 중심이며, 흔한 표현이나 추천 표현이 있어도 자연스러운 후기일 수 있습니다."),
        ("final_label", "SUSPICIOUS", "P_text가 학습해야 할 조작/부자연 패턴이 명확하다고 판단. 리뷰 신뢰도를 떨어뜨릴 수 있는 명확한 반복/복사/구조적 정보/직접 행동유도 패턴을 문맥에서 확인합니다."),
        ("final_label", "UNCERTAIN", "정상/의심을 문맥만으로 확정하기 어려워 학습에서 보류할 리뷰. 애매하면 억지로 SUSPICIOUS로 라벨링하지 않습니다."),
        ("final_label", "INVALID", "내용 없음, 깨진 값, 리뷰가 아닌 내용 등 리뷰 학습 데이터 자체로 사용하기 어려운 경우입니다."),
        ("confirmed_patterns", "입력 방법", "후보 규칙을 복사하지 말고 사람이 최종 확인한 패턴만 입력합니다. 여러 코드는 pipe(|)로 구분합니다. 예: REPETITION|PROMOTIONAL_CTA. 확인한 패턴이 없으면 빈칸으로 둡니다."),
        ("confirmed_patterns", "REPETITION", "동일하거나 매우 유사한 리뷰 문구가 반복됨. 흔한 짧은 표현이나 합리적 이유가 있는 반복은 주의해서 판단합니다."),
        ("confirmed_patterns", "PRODUCT_COPY", "실제 사용 후기보다 상품 설명/스펙을 복사한 듯한 내용입니다."),
        ("confirmed_patterns", "STRUCTURED_INFO", "후기보다 정보 항목을 정리한 형식이 강하게 나타남. 예: 장점/단점/총평, 번호 항목, 체크리스트. 형식만으로 최종 라벨을 확정하지 않습니다."),
        ("confirmed_patterns", "PROMOTIONAL_CTA", "독자에게 구매/사용을 직접 유도하는 표현. 예: 꼭 사세요, 구매하세요, 꼭 써보세요. 단순한 '추천할 만해요' 같은 경험 표현과 구분합니다."),
        ("HARD_NEGATIVE 후보", "최종 위험 패턴 아님", "SNS/광고/추천 단어가 있어도 개인의 실제 구매·사용 경험이 분명한 경우, 모델이 광고성 리뷰로 오판하지 않도록 검토하는 후보입니다. 사람이 실제 문맥을 보고 NORMAL 여부를 판단합니다. confirmed_patterns에는 입력하지 않습니다."),
        ("candidate_reasons", "힌트일 뿐", "REPETITION이나 PROMOTIONAL_CTA 후보라고 SUSPICIOUS로 확정하지 않습니다. HARD_NEGATIVE 후보라고 NORMAL로 확정하지 않습니다. final_label과 confirmed_patterns를 자동으로 채우지 않습니다."),
        ("use_for_training", "YES / NO", "사람이 최종 검토 후 학습 사용 여부를 직접 결정합니다. final_label만으로 자동 결정하지 않습니다. UNCERTAIN은 보통 NO로 판단할 수 있지만 자동 입력하지 않으며 INVALID도 자동 입력하지 않습니다."),
        ("reviewer_note", "자유 입력", "판단 근거나 애매한 부분을 자유롭게 기록합니다. 초기값은 빈칸입니다."),
        ("Excel 사용", "빈칸 / 긴 본문", "노란 final_label 칸은 아직 입력되지 않은 행입니다. 긴 본문은 수식 입력줄에서 확인하거나 필요할 때 행 높이를 조절하세요. 드롭다운 검증은 붙여넣기로 우회될 수 있으므로 최종 값을 확인하세요."),
        ("원본 보존", "로컬 검토 파일", "후보 정보와 행 순서를 보존하고 수동 입력 네 칸은 모두 빈칸으로 시작합니다. 원본 후보 파일은 수정하지 않습니다. 리뷰 원문을 포함한 outputs/ 파일은 Git에 올리지 않습니다."),
    ]
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    for column, width in (("A", 24), ("B", 28), ("C", 100)):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 66
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    sheet.row_dimensions[1].height = 24


def main():
    parser = argparse.ArgumentParser(description="후보 Excel을 빈 수동 라벨링 칸과 가이드가 있는 workbook으로 변환합니다.")
    parser.add_argument("excel_path", type=Path, nargs="?",
                        default=PROJECT_ROOT / "outputs" / "ptext_candidates.xlsx",
                        help="입력 후보 Excel (기본: outputs/ptext_candidates.xlsx)")
    args = parser.parse_args()
    source = args.excel_path.expanduser().resolve()
    output = PROJECT_ROOT / "outputs" / "ptext_labeling_workbook.xlsx"
    workbook = None
    temporary_path = None
    try:
        if source == output.resolve():
            raise ValueError("입력과 출력 경로가 같습니다. 원본 보호를 위해 중단합니다.")
        if output.exists():
            raise ValueError("출력 파일이 이미 있습니다. 기존 수동 라벨을 보호하기 위해 덮어쓰지 않습니다. 파일이 열려 있다면 닫고, 기존 결과를 다른 이름으로 보관한 뒤 다시 실행해주세요.")
        records = read_candidates(source)
        workbook = Workbook()
        build_labeling_sheet(workbook, records)
        build_guide_sheet(workbook)
        output.parent.mkdir(parents=True, exist_ok=True)
        # A unique path avoids touching temporary files left by another run.
        with tempfile.NamedTemporaryFile(
            mode="w+b", dir=output.parent, prefix=output.stem + ".",
            suffix=".tmp.xlsx", delete=False,
        ) as stream:
            temporary_path = Path(stream.name)
            workbook.save(stream)
            stream.flush()
            os.fsync(stream.fileno())
        if os.name == "nt":
            # Windows rename is atomic and fails if the destination exists,
            # including a file created after the initial existence check.
            temporary_path.rename(output)
        else:
            # POSIX rename can overwrite: publish atomically with an exclusive
            # hard link instead, then remove this run's temporary name below.
            os.link(temporary_path, output)
        print(f"수동 라벨링 workbook 저장: {output}\n후보 행 수: {len(records)}\n수동 입력 칸은 모두 빈칸입니다.")
    except PermissionError:
        print("오류: 파일 접근 또는 저장이 불가능합니다. 입력/출력 파일의 권한을 확인하고 Excel에서 열려 있다면 닫아주세요.", file=sys.stderr)
        return 1
    except (BadZipFile, InvalidFileException):
        print("오류: 읽을 수 있는 Excel 파일이 아닙니다. 정상적인 .xlsx 파일인지 확인해주세요.", file=sys.stderr)
        return 1
    except Exception as exc:
        print(f"오류: 라벨링 workbook을 만들지 못했습니다. {exc}", file=sys.stderr)
        return 1
    finally:
        if temporary_path is not None:
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError as exc:
                print(f"경고: 이번 실행의 임시 파일을 정리하지 못했습니다: {temporary_path} ({exc})", file=sys.stderr)
        if workbook is not None:
            workbook.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
