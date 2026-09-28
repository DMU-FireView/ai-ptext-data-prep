"""Use only synthetic workbooks under tmp_path, never real outputs/ files."""

from pathlib import Path
import subprocess
import sys

from openpyxl import Workbook, load_workbook
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import prepare_ptext_human_review as review


HEADERS = (
    "source_excel_row", "platform", "product_id", "review_id", "content",
    "rating", "candidate_reasons", "matched_features",
)
ROWS = (
    (12, "shop", 123, 456, "  원문\n공백과 이모지 🙂 유지  ", 4.5,
     "PRODUCT_DESCRIPTION_REPLACEMENT", "성분 설명"),
    (13, "shop", 123.5, 456.5, "=원문 문자열", 3, "HARD_NEGATIVE", "직접 경험"),
    (None,) * 8,
    (14, "shop", "00123", "00456", "같은 후보라도 직접 판단", None,
     "REPETITION|PROMOTIONAL_CTA", "반복; 추천"),
)


def make_source(path, headers=HEADERS, sheet_name="candidates", rows=ROWS):
    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = sheet_name
        sheet.append(list(headers))
        for record in rows:
            sheet.append(record)
        if rows:
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            sheet["C2"].number_format = "000000"
        workbook.save(path)
    finally:
        workbook.close()
    return path


@pytest.fixture
def source(tmp_path):
    return make_source(tmp_path / "source.xlsx")


def assert_review_workbook(workbook, expected_rows):
    assert workbook.sheetnames == ["candidates"]
    sheet = workbook["candidates"]
    assert tuple(cell.value for cell in sheet[1]) == (*HEADERS, "human_label", "review_note")
    assert sheet.max_column == 10
    assert sheet.max_row == expected_rows + 1
    for row in sheet.iter_rows(min_row=2, max_row=expected_rows + 1, min_col=9, max_col=10):
        assert all(cell.value is None for cell in row)
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == f"A1:J{expected_rows + 1}"
    validations = list(sheet.data_validations.dataValidation)
    assert len(validations) == 1
    validation = validations[0]
    assert validation.type == "list"
    assert validation.formula1 == '"SUSPICIOUS,NORMAL,UNCERTAIN"'
    assert [cell_range.bounds for cell_range in validation.sqref.ranges] == [
        (9, 2, 9, max(2, expected_rows + 1))
    ]
    assert validation.allow_blank
    assert validation.showDropDown is False
    assert validation.showErrorMessage and validation.errorStyle == "stop"
    for row_number in range(2, expected_rows + 2):
        assert sheet.cell(row_number, 5).alignment.wrap_text
        assert sheet.cell(row_number, 8).alignment.wrap_text


def test_prepare_preserves_cells_types_order_and_empty_manual_fields(source, tmp_path):
    before = source.read_bytes()
    original = load_workbook(source, data_only=False)
    workbook, count = review.prepare_workbook(source)
    try:
        assert count == len(ROWS)
        assert_review_workbook(workbook, len(ROWS))
        expected = original["candidates"]
        actual = workbook["candidates"]
        for source_row, target_row in zip(expected.iter_rows(), actual.iter_rows(max_col=8)):
            for left, right in zip(source_row, target_row):
                assert right.value == left.value
                assert type(right.value) is type(left.value)
                assert right.data_type == left.data_type
                assert right.number_format == left.number_format
        for coordinate in ("A2", "C2", "D2", "C3", "D3", "F2"):
            assert isinstance(actual[coordinate].value, (int, float))
        assert actual["C5"].value == "00123"
        assert actual["D5"].value == "00456"
        assert actual["E3"].data_type == "s"
        # Also verify persistence, not just in-memory validation/style objects.
        saved = tmp_path / "roundtrip.xlsx"
        workbook.save(saved)
        reopened = load_workbook(saved, data_only=False)
        try:
            assert_review_workbook(reopened, len(ROWS))
            assert list(reopened["candidates"].iter_rows(max_col=8, values_only=True)) == list(
                expected.iter_rows(values_only=True)
            )
        finally:
            reopened.close()
    finally:
        original.close()
        workbook.close()
    assert source.read_bytes() == before


def test_header_only_input(tmp_path):
    path = make_source(tmp_path / "empty.xlsx", rows=())
    workbook, count = review.prepare_workbook(path)
    try:
        assert count == 0
        assert_review_workbook(workbook, 0)
    finally:
        workbook.close()


def test_missing_file(tmp_path):
    with pytest.raises(ValueError, match="입력 파일이 없습니다"):
        review.prepare_workbook(tmp_path / "missing.xlsx")


def test_missing_candidates_sheet(tmp_path):
    source = make_source(tmp_path / "wrong_sheet.xlsx", sheet_name="other")
    with pytest.raises(ValueError, match="candidates 시트가 없습니다"):
        review.prepare_workbook(source)


@pytest.mark.parametrize("headers", [
    ("wrong_name", *HEADERS[1:]),
    (HEADERS[1], HEADERS[0], *HEADERS[2:]),
    HEADERS[:-1],
    (*HEADERS, "unexpected_column"),
])
def test_invalid_column_names_or_order(tmp_path, headers):
    source = make_source(tmp_path / "bad_header.xlsx", headers=headers)
    before = source.read_bytes()
    with pytest.raises(ValueError, match="헤더.*순서가 일치해야"):
        review.prepare_workbook(source)
    assert source.read_bytes() == before


def run_cli(source, output, tmp_path):
    return subprocess.run(
        [sys.executable, str(Path(review.__file__).resolve()), str(source), "--output", str(output)],
        cwd=tmp_path, capture_output=True, timeout=30, check=False,
    )


def test_cli_creates_output_in_new_parent_directory(source, tmp_path):
    before = source.read_bytes()
    output = tmp_path / "new_folder" / "review.xlsx"
    completed = run_cli(source, output, tmp_path)
    assert completed.returncode == 0, completed.stderr
    workbook = load_workbook(output, data_only=False)
    try:
        assert_review_workbook(workbook, len(ROWS))
    finally:
        workbook.close()
    assert source.read_bytes() == before
    assert list(output.parent.iterdir()) == [output]


def test_cli_refuses_existing_output(source, tmp_path):
    output = make_source(tmp_path / "existing.xlsx")
    workbook = load_workbook(output)
    workbook["candidates"]["I2"] = "NORMAL"
    workbook.save(output)
    workbook.close()
    before_output = output.read_bytes()
    before_source = source.read_bytes()
    completed = run_cli(source, output, tmp_path)
    assert completed.returncode != 0
    assert output.read_bytes() == before_output
    assert source.read_bytes() == before_source
    assert not list(tmp_path.glob("*.tmp.xlsx"))


def test_cli_refuses_same_input_output(source, tmp_path):
    before = source.read_bytes()
    completed = run_cli(source, source, tmp_path)
    assert completed.returncode != 0
    assert source.read_bytes() == before
    assert not list(tmp_path.glob("*.tmp.xlsx"))
