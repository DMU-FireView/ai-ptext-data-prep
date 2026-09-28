"""V2 regressions: all workbook writes remain inside pytest tmp_path."""

from collections import Counter
from copy import deepcopy
from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import build_ptext_training_dataset as v1
import build_ptext_training_dataset_v2 as v2
from generate_ptext_synthetic_v2 import generate_v2, starting_key, TYPES
from test_build_ptext_training_dataset import inputs


def test_normal_dedup_priority_seed_and_exact_values():
    rows = [dict(source_excel_row=i, content=text, label="NORMAL", label_origin=origin)
            for i, text, origin in [(2, "동일 원문", "sampled_non_candidate"),
                                    (3, "동일 원문", "human_candidate_review"),
                                    (4, "다른 원문", "sampled_non_candidate"),
                                    (5, "다른 원문", "sampled_non_candidate"),
                                    (6, "동일 원문 ", "sampled_non_candidate")]]
    before = deepcopy(rows)
    result, removed = v2.deduplicate_normal(rows)
    assert len(result) == 3 and len(removed) == 2
    assert next(r for r in result if r["content"] == "동일 원문")["source_excel_row"] == 3
    assert any(r["content"] == "동일 원문 " for r in result)
    assert (result, removed) == v2.deduplicate_normal(list(reversed(rows)))
    assert rows == before


@pytest.fixture(scope="module")
def parents():
    return [dict(platform=("ohouse", "oliveyoung", "elevenst", "musinsa")[i % 4],
                 product_id=i + .5, review_id=i + 100, source_excel_row=i + 2,
                 content=("청소솔 손잡이와 크림 보습에 대한 실제 사용 기록 " * 12)[:60 + i * 3])
            for i in range(70)]


def test_generation_diversity_lengths_and_determinism(parents):
    before = deepcopy(parents)
    rows = generate_v2(parents)
    assert rows == generate_v2(parents)
    assert rows != generate_v2(parents, seed=43)
    assert Counter(r["synthetic_type"] for r in rows) == {kind: 250 for kind in TYPES}
    assert len({r["content"] for r in rows}) == 1250
    source = {p["source_excel_row"]: p for p in parents}
    for row in rows:
        parent = source[row["parent_source_excel_row"]]
        assert .7 <= len(row["content"]) / len(parent["content"]) <= 1.3
        assert row["platform"] == parent["platform"]
        assert row["product_id"] == parent["product_id"]
        assert row["parent_review_id"] == parent["review_id"]
        assert row["platform"] not in row["content"]
    for kind in TYPES:
        members = [r for r in rows if r["synthetic_type"] == kind]
        for size in (3, 4):
            assert max(Counter(starting_key(r["content"], size) for r in members).values()) / 250 < .05
    assert max(Counter(r["synthetic_family_id"] for r in rows if r["synthetic_type"] == "COPY_VARIATION").values()) == 2
    assert parents == before


@pytest.fixture
def built(inputs, tmp_path):
    old_rows, _, _ = v1.build_dataset(*inputs, sample_count=24, synthetic_count=10)
    # Deliberate weak duplicate of human-reviewed NORMAL.
    duplicate = dict(old_rows[0], source_excel_row=22, review_id=120,
                     product_id=20.5, label_origin="sampled_non_candidate")
    old_rows.append(duplicate)
    path = tmp_path / "v1.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "dataset"
    sheet.append(v1.COLUMNS)
    for row in old_rows:
        sheet.append([row.get(column) for column in v1.COLUMNS])
    workbook.save(path)
    workbook.close()
    reports = [tmp_path / "v1_summary.md", tmp_path / "v1_audit.md"]
    for report in reports:
        report.write_text("preserve v1", encoding="utf-8")
    protected = {p: p.read_bytes() for p in [path, *inputs, *reports]}
    bundle = v2.build_v2(path, *inputs, synthetic_count=30)
    return bundle, protected, path


def test_build_preserves_v1_sources_labels_and_has_no_leakage(built, inputs, tmp_path):
    (rows, summary, audit, metadata), protected, path = built
    assert all(audit["mandatory"].values())
    assert metadata["removed_exact_normal"] == 1
    assert metadata["v1_synthetic_discarded"] == 10
    assert audit["exact"]["all"]["groups"] == 0
    assert audit["copy_family_max"] == 2
    assert not any(r["label"] == "UNCERTAIN" for r in rows)
    assert (rows, summary, audit, metadata) == v2.build_v2(path, *inputs, synthetic_count=30)
    output, report, audit_path = (tmp_path / n for n in ("v2.xlsx", "v2_summary.md", "v2_audit.md"))
    v2.save_bundle(rows, summary, audit, metadata, {}, output, report, audit_path)
    book = load_workbook(output, read_only=True)
    try:
        assert book.sheetnames == ["dataset", "summary"]
        saved = list(book["dataset"].values)
        assert saved[0] == v1.COLUMNS
        assert len(saved) == len(rows) + 1
        assert isinstance(saved[1][1], float)
        assert isinstance(saved[1][2], int)
    finally:
        book.close()
    assert "실제 positive recall" in audit_path.read_text(encoding="utf-8")
    assert all(p.read_bytes() == original for p, original in protected.items())
    before = output.read_bytes()
    with pytest.raises(ValueError):
        v2.save_bundle(rows, summary, audit, metadata, {}, output, report, audit_path)
    assert output.read_bytes() == before


def test_audit_detects_length_prefix_and_word_bias(built):
    (rows, _, _, _), _, _ = built
    changed = deepcopy(rows)
    for i, row in enumerate(changed):
        if row["label_origin"] == "synthetic":
            row["content"] = "동일 시작 표현 선택 구성 " * 100 + str(i)
    audit = v2.quality_audit(changed)
    assert not any(audit["targets"].values())
    assert audit["content_word_warnings"]


def test_leakage_and_conflict_gate_blocks_publication(built, tmp_path):
    (rows, summary, _, metadata), _, _ = built
    changed = deepcopy(rows)
    synthetic = next(r for r in changed if r["label_origin"] == "synthetic")
    parent = next(r for r in changed if r.get("source_excel_row") == synthetic["parent_source_excel_row"])
    synthetic["split"] = "test" if parent["split"] != "test" else "train"
    synthetic["content"] = parent["content"]
    audit = v2.quality_audit(changed)
    assert not audit["mandatory"]["label_conflict_zero"]
    assert not audit["mandatory"]["group_leakage_zero"]
    assert not audit["mandatory"]["parent_leakage_zero"]
    assert not audit["mandatory"]["cross_split_exact_zero"]
    paths = [tmp_path / name for name in ("bad.xlsx", "bad.md", "bad_audit.md")]
    with pytest.raises(ValueError, match="quality gate"):
        v2.save_bundle(changed, summary, audit, metadata, {}, *paths)
    assert not any(path.exists() for path in paths)


def test_target_failures_are_visible_in_report(built):
    (rows, summary, audit, metadata), _, _ = built
    changed = deepcopy(audit)
    changed["targets"]["mean_length_gap_within_20pct"] = False
    assert "mean_length_gap_within_20pct | FAIL" in v2.report_text(summary, changed, metadata, {})


def test_no_eligible_parent_and_empty_generation():
    assert generate_v2([], 0) == []
    with pytest.raises(ValueError):
        generate_v2([], 5)
